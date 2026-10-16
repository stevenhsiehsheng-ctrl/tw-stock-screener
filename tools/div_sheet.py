"""產生付費商品「存股股息追蹤試算表」的 .xlsx（要用 Google 試算表開：資料靠 IMPORTDATA 每天自動更新）。

資料來源是網站每天收盤後產生的 site/data/div_summary.csv、dividends.csv（screener/divpages.py）。
用法：python tools/div_sheet.py --out 存股股息追蹤試算表.xlsx [--base https://…/tw-stock-screener/]
"""
from __future__ import annotations

import argparse

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

BASE = "https://stevenhsiehsheng-ctrl.github.io/tw-stock-screener/"
ROWS = 50            # 持股最多幾列
FIRST = 6            # 持股第一列
EXAMPLE = [("0056", 5000, 33.5), ("00878", 8000, 21.2), ("00919", 6000, 23.8), ("2330", 100, 980)]
DISCLAIMER = "本試算表僅整理證交所、櫃買中心公開資訊，不構成任何投資建議；資料可能有誤或延遲，以官方公告為準，投資盈虧請自行負責。"

HEAD = PatternFill("solid", fgColor="2F5BD3")
INPUT = PatternFill("solid", fgColor="FFF4CC")
WHITE = Font(color="FFFFFF", bold=True)
BOLD = Font(bold=True)


def build(base: str, out: str) -> None:
    wb = Workbook()

    # ---------------------------------------------------------------- 我的持股
    ws = wb.active
    ws.title = "我的持股"
    ws["A1"] = "存股股息追蹤試算表"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = "黃色欄位填你的持股（代號、股數、平均成本），其他欄位會自動算；資料每天收盤後自動更新。"
    last = FIRST + ROWS - 1
    rng = lambda c: f"{c}{FIRST}:{c}{last}"   # noqa: E731
    ws["A3"], ws["B3"] = "總市值", f"=SUM({rng('N')})"
    ws["C3"], ws["D3"] = "預估年股息", f"=SUM({rng('H')})"
    ws["E3"], ws["F3"] = "平均每月", "=D3/12"
    ws["G3"], ws["H3"] = "整體殖利率", "=IFERROR(D3/B3,0)"
    ws["I3"], ws["J3"] = "30 天內除息", f'=COUNTIFS({rng("K")},">="&TODAY(),{rng("K")},"<="&TODAY()+30)'
    for c in "ACEGI":
        ws[f"{c}3"].font = BOLD
    for c, f in (("B", "#,##0"), ("D", "#,##0"), ("F", "#,##0"), ("H", "0.00%")):
        ws[f"{c}3"].number_format = f
        ws[f"{c}3"].font = Font(bold=True, size=13)

    heads = ["代號", "持有股數", "平均成本（選填）", "名稱", "收盤價", "近 12 月每股配息", "配息次數（近 12 月）",
             "預估年股息", "現價殖利率", "成本殖利率", "下次除息日", "下次每股配息", "下次預估股息", "市值"]
    for i, h in enumerate(heads, 1):
        c = ws.cell(FIRST - 1, i, h)
        c.fill, c.font = HEAD, WHITE
        c.alignment = Alignment(wrap_text=True, vertical="center")
    look = lambda r, col: f'IFERROR(VLOOKUP("#"&TRIM($A{r}),資料_摘要!$A:$M,{col},FALSE),"")'  # noqa: E731
    for r in range(FIRST, last + 1):
        for c in "ABC":
            ws[f"{c}{r}"].fill = INPUT
        ws[f"A{r}"].number_format = "@"
        ws[f"D{r}"] = f'=IF($A{r}="","",IFERROR(VLOOKUP("#"&TRIM($A{r}),資料_摘要!$A:$M,3,FALSE),"查無此代號"))'
        ws[f"E{r}"] = f'=IF($A{r}="","",{look(r, 4)})'
        ws[f"F{r}"] = f'=IF($A{r}="","",{look(r, 6)})'
        ws[f"G{r}"] = f'=IF($A{r}="","",{look(r, 7)})'
        ws[f"H{r}"] = f'=IF($A{r}="","",IFERROR(F{r}*B{r},0))'
        ws[f"I{r}"] = f'=IF($A{r}="","",IFERROR(F{r}/E{r},""))'
        ws[f"J{r}"] = f'=IF(OR($A{r}="",C{r}="",C{r}=0),"",IFERROR(F{r}/C{r},""))'
        ws[f"K{r}"] = f'=IF($A{r}="","",IFERROR(DATEVALUE({look(r, 11)}),{look(r, 11)}))'
        ws[f"L{r}"] = f'=IF($A{r}="","",IF(K{r}="","",IF({look(r, 12)}="","待公告",{look(r, 12)})))'
        ws[f"M{r}"] = f'=IF(ISNUMBER(L{r}),L{r}*B{r},"")'
        ws[f"N{r}"] = f'=IF($A{r}="","",IFERROR(E{r}*B{r},0))'
        for c, f in (("B", "#,##0"), ("C", "0.00"), ("E", "0.00"), ("F", "0.000"), ("H", "#,##0"), ("I", "0.00%"),
                     ("J", "0.00%"), ("K", "yyyy-mm-dd"), ("L", "0.000"), ("M", "#,##0"), ("N", "#,##0")):
            ws[f"{c}{r}"].number_format = f
    for (code, sh, cost), r in zip(EXAMPLE, range(FIRST, last + 1)):
        ws[f"A{r}"], ws[f"B{r}"], ws[f"C{r}"] = code, sh, cost
    dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"B{FIRST}:C{last}")
    for col, w in zip("ABCDEFGHIJKLMN", (10, 11, 12, 18, 10, 12, 11, 12, 10, 10, 12, 12, 12, 12)):
        ws.column_dimensions[col].width = w
    ws.row_dimensions[FIRST - 1].height = 34
    ws.freeze_panes = f"B{FIRST}"
    ws[f"A{last + 2}"] = DISCLAIMER
    ws[f"A{last + 2}"].font = Font(size=9, color="777777")

    # ---------------------------------------------------------------- 每月股息
    m = wb.create_sheet("每月股息")
    m["A1"] = "今年每月領多少（依除息日所在月份；實際入帳通常晚約 3～4 週）"
    m["A1"].font = Font(bold=True, size=14)
    m["A2"] = "已除息的照實際金額；還沒除息的只算已公告金額，待公告的先算 0。股數用「我的持股」現在填的。"
    hr = 4
    m.cell(hr, 1, "代號").fill, m.cell(hr, 2, "股數").fill = HEAD, HEAD
    m.cell(hr, 1).font = m.cell(hr, 2).font = WHITE
    for k in range(12):
        c = m.cell(hr, 3 + k, f'=TEXT(DATE(YEAR(TODAY()),{k + 1},1),"yyyy-mm")')
        c.fill, c.font = HEAD, WHITE
        m.cell(hr - 1, 3 + k, f'="m"&TEXT(DATE(YEAR(TODAY()),{k + 1},1),"yyyymm")').font = Font(color="BBBBBB", size=8)
    m.cell(hr, 15, "全年").fill, m.cell(hr, 15).font = HEAD, WHITE
    for i in range(ROWS):
        r, src = hr + 1 + i, FIRST + i
        m[f"A{r}"] = f'=IF(我的持股!A{src}="","",我的持股!A{src})'
        m[f"B{r}"] = f'=IF(我的持股!A{src}="","",我的持股!B{src})'
        for k in range(12):
            col = m.cell(r, 3 + k).column_letter
            m[f"{col}{r}"] = (f'=IF($A{r}="","",SUMIFS(資料_配息!$E:$E,資料_配息!$A:$A,"#"&TRIM($A{r}),'
                              f'資料_配息!$G:$G,{col}${hr - 1})*$B{r})')
            m[f"{col}{r}"].number_format = "#,##0"
        m[f"O{r}"] = f'=IF($A{r}="","",SUM(C{r}:N{r}))'
        m[f"O{r}"].number_format = "#,##0"
    tot = hr + ROWS + 1
    m[f"A{tot}"] = "合計"
    m[f"A{tot}"].font = BOLD
    for k in range(13):
        col = m.cell(tot, 3 + k).column_letter
        m[f"{col}{tot}"] = f"=SUM({col}{hr + 1}:{col}{tot - 1})"
        m[f"{col}{tot}"].number_format = "#,##0"
        m[f"{col}{tot}"].font = BOLD
    ch = BarChart()
    ch.title, ch.y_axis.title, ch.legend = "今年每月股息（元）", "元", None
    ch.add_data(Reference(m, min_col=3, max_col=14, min_row=tot), from_rows=True)
    ch.set_categories(Reference(m, min_col=3, max_col=14, min_row=hr))
    ch.width, ch.height = 22, 9
    m.add_chart(ch, "Q4")
    m.freeze_panes = f"C{hr + 1}"

    # ---------------------------------------------------------------- 使用說明
    h = wb.create_sheet("使用說明")
    lines = [
        ("存股股息追蹤試算表：使用說明", True),
        ("", False),
        ("1. 這份檔案要用 Google 試算表開（Excel 不支援自動更新）。", False),
        ("   打開 Google 雲端硬碟 → 新增 → 檔案上傳 → 上傳後點兩下，用 Google 試算表開啟。", False),
        ("   若跳出「允許存取外部資料」，請按允許，資料才會載入。", False),
        ("2. 到「我的持股」，把黃色欄位換成你自己的股票：代號、持有股數、平均成本（選填）。", False),
        ("   上市櫃個股和 ETF 都可以，例如 0056、00878、2330。範例那幾列直接改掉或刪掉就好。", False),
        ("3. 其他欄位會自動算：近 12 月配息、殖利率、下次除息日、預估年股息、每月領多少。", False),
        ("4. 資料每天收盤後自動更新，不用做任何事。試算表通常每小時會重新抓一次。", False),
        ("", False),
        ("常見問題", True),
        ("Q：顯示「查無此代號」？ A：確認代號正確；只配股不配現金、或兩年內沒配過息的股票不會出現。", False),
        ("Q：下次每股配息顯示「待公告」？ A：公司已公告除息日但還沒公告金額，公告後會自動補上。", False),
        ("Q：殖利率怎麼算？ A：近 12 個月已除息的現金股利合計 ÷ 最近收盤價，是過去的數字，不代表未來。", False),
        ("Q：股息什麼時候入帳？ A：通常在除息日後 3～4 週的發放日，「每月股息」頁是依除息日分月。", False),
        ("", False),
        (DISCLAIMER, False),
    ]
    for i, (t, b) in enumerate(lines, 1):
        h.cell(i, 1, t).font = Font(bold=b, size=14 if b else 11)
    h.column_dimensions["A"].width = 110

    # ---------------------------------------------------------------- 資料（自動）
    for name, f in (("資料_摘要", "div_summary.csv"), ("資料_配息", "dividends.csv")):
        d = wb.create_sheet(name)
        d["A1"] = f'=IMPORTDATA("{base}data/{f}")'
        d.sheet_properties.tabColor = "BBBBBB"
    wb.move_sheet("使用說明", offset=-2)
    wb.active = 1
    wb.save(out)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default="存股股息追蹤試算表.xlsx")
    ap.add_argument("--base", default=BASE)
    a = ap.parse_args()
    build(a.base, a.out)
    print("寫好了：", a.out)
